import io
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_text(filename, content):
    if not content or len(content) > 3 * 1024 * 1024:
        raise ValueError("Upload a non-empty file up to 3 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        text = content.decode("utf-8-sig")
    elif suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        if len(reader.pages) > 40:
            raise ValueError("PDFs are limited to 40 pages.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix in {".docx", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                raise ValueError("Expanded document is too large.")
        if suffix == ".docx":
            doc = Document(io.BytesIO(content))
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [
                    " | ".join(c.text for c in row.cells)
                    for table in doc.tables
                    for row in table.rows
                ]
            )
        else:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            text = "\n".join(
                f"{sheet.title}: " + " | ".join(str(cell) for cell in row if cell is not None)
                for sheet in book
                for row in sheet.iter_rows(values_only=True)
            )
            book.close()
    else:
        raise ValueError("Supported files: PDF, DOCX, XLSX, TXT, Markdown and CSV.")
    text = text.strip()
    if not text:
        raise ValueError(
            "No readable text found. Scanned PDFs need OCR; upload a text version for now."
        )
    if len(text) > 80000:
        raise ValueError("Extracted text is limited to 80,000 characters per file.")
    return text


def resolve_event(message, events, selected):
    matches = [
        event
        for event in events
        if re.search(r"(?<!\w)" + re.escape(event["name"]) + r"(?!\w)", message, re.IGNORECASE)
    ]
    matches = [
        event
        for event in matches
        if not any(
            event["name"].casefold() != other["name"].casefold()
            and event["name"].casefold() in other["name"].casefold()
            for other in matches
        )
    ]
    if len(matches) > 1 and len({e["name"].casefold() for e in matches}) == 1:
        return next((e for e in matches if e["id"] == selected), None)
    if len(matches) > 1:
        raise ValueError("You mentioned multiple events. Choose one for this conversation.")
    if matches:
        return matches[0]
    return next((event for event in events if event["id"] == selected), None)


def retrieve(event_id, query, documents):
    words = set(re.findall(r"\w+", query.lower()))
    chunks = []
    for document in documents:
        if document["eventId"] != event_id:
            continue
        for index in range(0, len(document["text"]), 1200):
            excerpt = document["text"][index : index + 1200]
            score = len(words & set(re.findall(r"\w+", excerpt.lower())))
            chunks.append(
                (
                    score,
                    {
                        "documentId": document["id"],
                        "name": document["name"],
                        "chunk": index // 1200,
                        "excerpt": excerpt,
                    },
                )
            )
    return [chunk for _, chunk in sorted(chunks, key=lambda row: row[0], reverse=True)[:6]]


def amount(text):
    match = re.search(
        r"(?:RM\s*([\d,]+(?:\.\d{1,2})?)\s*(k)?\b|([\d,]+(?:\.\d{1,2})?)\s*(k)?\s*RM\b)",
        text,
        re.IGNORECASE,
    )
    if not match:
        return None
    number = match.group(1) or match.group(3)
    multiplier = 1000 if match.group(2) or match.group(4) else 1
    return round(float(number.replace(",", "")) * multiplier * 100)


def facts_from_text(text):
    result = {}
    labelled_budget = re.search(
        r"(?:budget|within|under|less than|spend up to|maximum spend|limit)[^\n.!?]{0,35}?(?:RM\s*[\d,]+(?:\.\d{1,2})?\s*k?\b|[\d,]+(?:\.\d{1,2})?\s*k?\s*RM\b)",
        text,
        re.IGNORECASE,
    )
    budget = amount(labelled_budget.group(0)) if labelled_budget else None
    if budget is not None:
        result["budgetSen"] = budget
    quantity = re.search(
        r"(\d+)\s+(?:more\s+|cotton\s+|custom\s+|printed\s+)*(?:tote(?:\s+bags)?|bags|items|units|(?:t-?)?shirts|shirts|people|persons?|attendees|pax)",
        text,
        re.IGNORECASE,
    )
    if quantity:
        result["quantity"] = int(quantity.group(1))
    budget_pattern = r"(?:budget|within|under|less than|spend up to|maximum spend|limit)[^\n.!?]{0,35}?(?:RM\s*[\d,]+(?:\.\d{1,2})?\s*k?\b|[\d,]+(?:\.\d{1,2})?\s*k?\s*RM\b)"
    total_budget = None
    unit_budget = None
    for match in re.finditer(budget_pattern, text, re.IGNORECASE):
        value = amount(match.group(0))
        if re.match(r"\s*(?:each|apiece|per\s+(?:item|unit|piece|(?:t-?)?shirt))\b", text[match.end():], re.IGNORECASE):
            unit_budget = value
        else:
            total_budget = value
    if unit_budget is not None:
        result["unitBudgetSen"] = unit_budget
        result.pop("budgetSen", None)
        if total_budget is not None:
            result["budgetSen"] = total_budget
            result["budgetSource"] = "total"
        elif quantity:
            result["budgetSen"] = unit_budget * result["quantity"]
            result["budgetSource"] = "unit"
    deadline = re.search(
        r"(?:deadline|deliver(?:y)?(?:\s+date)?|arrive(?:\s+by)?|need(?:ed)?\s+by|event\s+date)[^\n]{0,30}?\b(20\d{2}-\d{2}-\d{2})\b",
        text,
        re.IGNORECASE,
    )
    if deadline:
        from datetime import date

        try:
            result["deadline"] = date.fromisoformat(deadline.group(1)).isoformat()
        except ValueError:
            pass
    natural_date = re.search(
        r"\b(?:by|deadline(?:\s+is)?|deliver(?:y)?(?:\s+by)?)\s+(\d{1,2})\s+(January|February|March|April|May|June|July|August|September|October|November|December)\s+(20\d{2})\b",
        text,
        re.IGNORECASE,
    )
    if natural_date:
        from datetime import date

        months = [
            "january",
            "february",
            "march",
            "april",
            "may",
            "june",
            "july",
            "august",
            "september",
            "october",
            "november",
            "december",
        ]
        try:
            result["deadline"] = date(
                int(natural_date.group(3)),
                months.index(natural_date.group(2).lower()) + 1,
                int(natural_date.group(1)),
            ).isoformat()
        except ValueError:
            pass
    location = re.search(
        r"(?:location|venue|delivery location)\s*:\s*([^\n]+)", text, re.IGNORECASE
    )
    if location:
        result["location"] = location.group(1).strip()[:200]
    else:
        conversational_location = re.search(
            r"\b(?:in|deliver(?:y)?\s+to|located\s+in)\s+([A-Z][A-Za-z]+(?:[ -][A-Z][A-Za-z]+){0,5})(?=[.!?,;\n]|$|\s+(?:by|before|within|for)\b)",
            text,
        )
        if conversational_location:
            result["location"] = conversational_location.group(1).strip()[:200]
    if "location" not in result and re.search(r"\b(?:we(?:'ll| will| can)?|I(?:'ll| will| can)?|prefer|arrange|for|by)\s+(?:self[- ]?)?(?:collect(?:ion)?|pick[- ]?up)\b|\b(?:self[- ]collection|collection only|pick[- ]?up only)\b", text, re.IGNORECASE):
        result["location"] = "Collection from supplier (address to be confirmed)"
    return result


def intake(message, event):
    facts = event.get("facts", {})
    overrides = facts_from_text(message)
    budget = overrides.get("budgetSen", facts.get("budgetSen"))
    return {
        "title": f"Partners for {event['name']}",
        "goal": message,
        "eventId": event["id"],
        "quantity": overrides.get("quantity", facts.get("quantity")),
        "budgetSen": budget,
        "unitBudgetSen": overrides.get("unitBudgetSen", facts.get("unitBudgetSen")),
        "targetSen": round(budget * 0.9) if budget else None,
        "deadline": overrides.get("deadline", facts.get("deadline")),
        "location": overrides.get("location", facts.get("location")),
    }
